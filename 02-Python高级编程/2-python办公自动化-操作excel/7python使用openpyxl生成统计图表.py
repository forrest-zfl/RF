from openpyxl import Workbook
from openpyxl.chart import BarChart,Reference
wb = Workbook(write_only=True)
sheet = wb.create_sheet("手机销量统计")
rows = [
    ('类别','G战区',"X战区"),
    ('小米',88,76),
    ('vivo',68,89),
    ('oppo',87,59)
]
# 向工作表中添加行
for row in rows:
    sheet.append(row)
# 创建图表对象
chart = BarChart()
chart.type = "col"
chart.style = 10
# 设置图表的标题
chart.title = "手机销量统计"
# 设置纵轴标题
chart.y_axis.title = "销量"
# 设置横轴标题
chart.x_axis.title = "手机品牌"
# 设置数据的范围
data = Reference(sheet, min_col=2, min_row=1, max_row=5, max_col=3)
# 设置分类的范围
cats = Reference(sheet, min_col=1, min_row=2, max_row=5)
# 给图表添加数据
chart.add_data(data,titles_from_data = True)
# 给图表设置分类
chart.set_categories(cats)
chart.shape = 4
# 将图表插入到指定的单元格中
sheet.add_chart(chart,"A10")
wb.save("手机销量统计.xlsx")